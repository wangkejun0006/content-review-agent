from __future__ import annotations

import hashlib
import io
import zipfile
from pathlib import Path
from typing import Protocol

import httpx
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from smart_growth_agent.config import Settings
from smart_growth_agent.domain.errors import ValidationError
from smart_growth_agent.schemas.document_parse import (
    DocumentBlock,
    DocumentParseReport,
    ParsedDocument,
    ParseNotice,
    SourceLocator,
)

NATIVE_SUFFIXES = {".txt", ".md", ".pdf", ".docx", ".xlsx"}
MINERU_SUFFIXES = NATIVE_SUFFIXES | {
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
    ".pptx",
}
MAX_DOCUMENT_CHARACTERS = 200_000
MAX_PDF_PAGES = 300
MAX_WORKBOOK_CELLS = 100_000
MAX_WORKBOOK_SHEETS = 20


class DocumentParser(Protocol):
    async def parse(self, *, filename: str, media_type: str, content: bytes) -> ParsedDocument: ...


class NativeDocumentParser:
    parser_name = "smart-growth-native"
    parser_version = "1"

    async def parse(
        self,
        *,
        filename: str,
        media_type: str,
        content: bytes,
    ) -> ParsedDocument:
        suffix = Path(filename).suffix.lower()
        if suffix not in NATIVE_SUFFIXES:
            raise ValidationError(
                f"本地解析器暂不支持 {suffix or '无扩展名'} 文件；"
                "图片、PPT 和扫描件需要配置 MinerU"
            )
        try:
            if suffix in {".txt", ".md"}:
                blocks, notices = self._parse_text(content)
            elif suffix == ".pdf":
                blocks, notices = self._parse_pdf(content)
            elif suffix == ".docx":
                blocks, notices = self._parse_docx(content)
            else:
                blocks, notices = self._parse_xlsx(content)
        except ValidationError:
            raise
        except Exception as exc:
            raise ValidationError(f"无法读取文件 {filename}：{exc}") from exc
        return _build_document(
            filename=filename,
            media_type=media_type,
            content=content,
            parser_name=self.parser_name,
            parser_version=self.parser_version,
            blocks=blocks,
            notices=notices,
        )

    @staticmethod
    def _parse_text(content: bytes) -> tuple[list[DocumentBlock], list[ParseNotice]]:
        text = content.decode("utf-8-sig").strip()
        blocks = (
            [
                DocumentBlock(
                    block_id="text-1",
                    block_type="text",
                    text=text,
                )
            ]
            if text
            else []
        )
        return blocks, []

    @staticmethod
    def _parse_pdf(content: bytes) -> tuple[list[DocumentBlock], list[ParseNotice]]:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValidationError(f"PDF 超过 {MAX_PDF_PAGES} 页，请拆分后再上传")
        blocks: list[DocumentBlock] = []
        empty_pages = 0
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                empty_pages += 1
                continue
            blocks.append(
                DocumentBlock(
                    block_id=f"page-{page_number}",
                    block_type="text",
                    text=text,
                    locator=SourceLocator(page_number=page_number),
                )
            )
        notices: list[ParseNotice] = []
        if empty_pages:
            notices.append(
                ParseNotice(
                    code="pdf_pages_without_text",
                    severity="warning",
                    message=(
                        f"{empty_pages} 个 PDF 页面没有可提取文本，可能是扫描页；"
                        "配置 MinerU 后才能完整识别"
                    ),
                )
            )
        return blocks, notices

    @staticmethod
    def _parse_docx(content: bytes) -> tuple[list[DocumentBlock], list[ParseNotice]]:
        document = Document(io.BytesIO(content))
        blocks: list[DocumentBlock] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if text:
                blocks.append(
                    DocumentBlock(
                        block_id=f"paragraph-{index}",
                        block_type="text",
                        text=text,
                    )
                )
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                values = [cell.text.strip() for cell in row.cells]
                text = " | ".join(value for value in values if value)
                if text:
                    blocks.append(
                        DocumentBlock(
                            block_id=f"table-{table_index}-row-{row_index}",
                            block_type="table_row",
                            text=text,
                            locator=SourceLocator(
                                table_index=table_index,
                                row_index=row_index,
                            ),
                        )
                    )
        notices: list[ParseNotice] = []
        inline_shapes = len(document.inline_shapes)
        if inline_shapes:
            notices.append(
                ParseNotice(
                    code="docx_images_not_analyzed",
                    severity="warning",
                    message=(
                        f"Word 包含 {inline_shapes} 个内嵌图片，本地解析器未分析图片文字；"
                        "建议使用 MinerU 复核"
                    ),
                )
            )
        return blocks, notices

    @staticmethod
    def _parse_xlsx(content: bytes) -> tuple[list[DocumentBlock], list[ParseNotice]]:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            packaged_media_count = len(
                [
                    name
                    for name in archive.namelist()
                    if name.startswith("xl/media/") and not name.endswith("/")
                ]
            )
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=False)
        blocks: list[DocumentBlock] = []
        notices: list[ParseNotice] = []
        visited_cells = 0
        image_count = 0
        formula_count = 0
        try:
            if len(workbook.worksheets) > MAX_WORKBOOK_SHEETS:
                raise ValidationError(
                    f"Excel 超过 {MAX_WORKBOOK_SHEETS} 个工作表，请删除无关工作表或拆分后上传"
                )
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                image_count += len(getattr(sheet, "_images", []))
                for row_index, row in enumerate(sheet.iter_rows(), start=1):
                    visited_cells += len(row)
                    if visited_cells > MAX_WORKBOOK_CELLS:
                        raise ValidationError(
                            f"Excel 超过 {MAX_WORKBOOK_CELLS} 个待扫描单元格，请拆分后再上传"
                        )
                    populated = [cell for cell in row if cell.value is not None]
                    if not populated:
                        continue
                    rendered: list[str] = []
                    for cell in populated:
                        value = str(cell.value).strip()
                        if value.startswith("="):
                            formula_count += 1
                        rendered.append(value)
                    blocks.append(
                        DocumentBlock(
                            block_id=f"sheet-{sheet_index}-row-{row_index}",
                            block_type="table_row",
                            text=" | ".join(rendered),
                            locator=SourceLocator(
                                sheet_name=sheet.title,
                                cell_range=f"{populated[0].coordinate}:{populated[-1].coordinate}",
                                row_index=row_index,
                            ),
                        )
                    )
        finally:
            workbook.close()
        image_count = max(image_count, packaged_media_count)
        if image_count:
            notices.append(
                ParseNotice(
                    code="xlsx_images_not_analyzed",
                    severity="warning",
                    message=(
                        f"Excel 包含 {image_count} 个嵌入图片，本地解析器未分析图片文字；"
                        "正式入库前应使用 MinerU 复核"
                    ),
                )
            )
        if formula_count:
            notices.append(
                ParseNotice(
                    code="xlsx_formulas_preserved",
                    severity="information",
                    message=f"Excel 中有 {formula_count} 个公式，当前保留公式原文而非计算结果",
                )
            )
        return blocks, notices


class MinerUHTTPParser:
    parser_name = "mineru-http"
    parser_version = "3-api"

    def __init__(self, *, base_url: str, timeout_seconds: float) -> None:
        self.base_url = base_url.rstrip("/")
        self.timeout_seconds = timeout_seconds

    async def healthcheck(self) -> dict:
        async with httpx.AsyncClient(timeout=self.timeout_seconds) as client:
            response = await client.get(f"{self.base_url}/health")
            response.raise_for_status()
            return response.json()

    async def parse(
        self,
        *,
        filename: str,
        media_type: str,
        content: bytes,
    ) -> ParsedDocument:
        suffix = Path(filename).suffix.lower()
        if suffix not in MINERU_SUFFIXES:
            raise ValidationError(f"MinerU 暂不支持 {suffix or '无扩展名'} 文件")
        try:
            async with httpx.AsyncClient(timeout=self.timeout_seconds) as client:
                response = await client.post(
                    f"{self.base_url}/file_parse",
                    files={"files": (filename, content, media_type)},
                    data={
                        "return_md": "true",
                        "response_format_zip": "true",
                        "return_original_file": "false",
                    },
                )
                response.raise_for_status()
        except httpx.HTTPError as exc:
            raise ValidationError(f"MinerU 解析服务不可用：{exc}") from exc
        markdown = self._read_markdown_zip(response.content)
        blocks = [
            DocumentBlock(
                block_id="mineru-markdown-1",
                block_type="markdown",
                text=markdown,
            )
        ]
        return _build_document(
            filename=filename,
            media_type=media_type,
            content=content,
            parser_name=self.parser_name,
            parser_version=self.parser_version,
            blocks=blocks,
            notices=[],
        )

    @staticmethod
    def _read_markdown_zip(content: bytes) -> str:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = sorted(name for name in archive.namelist() if name.endswith(".md"))
                if not names:
                    raise ValidationError("MinerU 返回结果中没有 Markdown 文件")
                return "\n\n".join(
                    archive.read(name).decode("utf-8").strip() for name in names
                ).strip()
        except zipfile.BadZipFile as exc:
            raise ValidationError("MinerU 返回了无法读取的结果包") from exc


class DocumentParserRouter:
    def __init__(
        self,
        *,
        native: NativeDocumentParser,
        mineru: MinerUHTTPParser | None,
    ) -> None:
        self.native = native
        self.mineru = mineru

    async def parse(
        self,
        *,
        filename: str,
        media_type: str,
        content: bytes,
    ) -> ParsedDocument:
        suffix = Path(filename).suffix.lower()
        if suffix not in NATIVE_SUFFIXES:
            if self.mineru is None:
                raise ValidationError(
                    "该文件需要 MinerU 解析，但尚未配置 SMART_GROWTH_MINERU_BASE_URL"
                )
            return await self.mineru.parse(
                filename=filename,
                media_type=media_type,
                content=content,
            )
        parsed = await self.native.parse(
            filename=filename,
            media_type=media_type,
            content=content,
        )
        needs_mineru = any(
            notice.code
            in {
                "pdf_pages_without_text",
                "docx_images_not_analyzed",
                "xlsx_images_not_analyzed",
            }
            for notice in parsed.report.notices
        )
        if needs_mineru and self.mineru is not None:
            try:
                return await self.mineru.parse(
                    filename=filename,
                    media_type=media_type,
                    content=content,
                )
            except ValidationError as exc:
                parsed.report.notices.append(
                    ParseNotice(
                        code="mineru_fallback_to_native",
                        severity="warning",
                        message=f"{exc}；已回退到本地解析结果，正式入库前仍需复核",
                    )
                )
                parsed.report.complete = False
        return parsed


def build_document_parser(settings: Settings) -> DocumentParserRouter:
    mineru = (
        MinerUHTTPParser(
            base_url=settings.mineru_base_url,
            timeout_seconds=settings.mineru_timeout_seconds,
        )
        if settings.mineru_base_url
        else None
    )
    return DocumentParserRouter(native=NativeDocumentParser(), mineru=mineru)


def _build_document(
    *,
    filename: str,
    media_type: str,
    content: bytes,
    parser_name: str,
    parser_version: str,
    blocks: list[DocumentBlock],
    notices: list[ParseNotice],
) -> ParsedDocument:
    character_count = sum(len(block.text) for block in blocks)
    if character_count > MAX_DOCUMENT_CHARACTERS:
        raise ValidationError(
            f"文件 {filename} 提取出超过 {MAX_DOCUMENT_CHARACTERS} 个字符，请拆分后再上传"
        )
    if not blocks:
        notices.append(
            ParseNotice(
                code="no_text_extracted",
                severity="blocker",
                message="没有提取到可供 AI 分析的内容",
            )
        )
    complete = not any(notice.severity in {"warning", "blocker"} for notice in notices)
    return ParsedDocument(
        report=DocumentParseReport(
            source_name=filename,
            media_type=media_type,
            checksum=hashlib.sha256(content).hexdigest(),
            parser_name=parser_name,
            parser_version=parser_version,
            character_count=character_count,
            block_count=len(blocks),
            complete=complete,
            notices=notices,
        ),
        blocks=blocks,
    )
