from __future__ import annotations

import io

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from sqlalchemy.exc import SQLAlchemyError

from smart_growth_agent.app.api.product_import import MIB, validate_product_upload
from smart_growth_agent.db.models import Brand
from smart_growth_agent.domain.enums import RegistryType
from smart_growth_agent.domain.errors import ConflictError, ValidationError
from smart_growth_agent.llm.fake import DevelopmentFakeLLMGateway
from smart_growth_agent.schemas.registry import RegistryItemCreate
from smart_growth_agent.services.document_parser import NativeDocumentParser
from smart_growth_agent.services.ingestion import IngestionDedupService
from smart_growth_agent.services.product_import import ProductImportService
from smart_growth_agent.services.registry import RegistryService


def test_product_import_extracts_excel_text() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "规格"
    sheet.append(["型号", "NOTE01"])
    sheet.append(["充电方式", "磁吸充电"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    text = ProductImportService.extract_text("NOTE01.xlsx", buffer.getvalue())
    assert "工作表：规格" in text
    assert "型号 | NOTE01" in text
    assert "单元格=A1:B1" in text


def test_product_upload_limits_are_explicit() -> None:
    validate_product_upload("spec.xlsx", 8 * MIB)
    with pytest.raises(ValidationError, match="旧版 .xls"):
        validate_product_upload("legacy.xls", 1024)
    with pytest.raises(ValidationError, match="超过 8 MB"):
        validate_product_upload("large.xlsx", 8 * MIB + 1)


@pytest.mark.asyncio
async def test_native_parser_preserves_docx_table_provenance() -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "型号"
    table.cell(0, 1).text = "NOTE01"
    table.cell(1, 0).text = "材质"
    table.cell(1, 1).text = "天然石"
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = await NativeDocumentParser().parse(
        filename="spec.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=buffer.getvalue(),
    )

    assert parsed.report.complete is True
    assert parsed.blocks[0].locator.table_index == 1
    assert parsed.blocks[0].locator.row_index == 1
    assert parsed.blocks[0].text == "型号 | NOTE01"


@pytest.mark.asyncio
async def test_native_parser_marks_textless_pdf_as_incomplete() -> None:
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)

    parsed = await NativeDocumentParser().parse(
        filename="scan.pdf",
        media_type="application/pdf",
        content=buffer.getvalue(),
    )

    assert parsed.report.complete is False
    assert {notice.code for notice in parsed.report.notices} == {
        "pdf_pages_without_text",
        "no_text_extracted",
    }


@pytest.mark.asyncio
async def test_native_parser_reports_xlsx_embedded_images() -> None:
    image_buffer = io.BytesIO()
    Image.new("RGB", (20, 20), color="blue").save(image_buffer, format="PNG")
    image_buffer.seek(0)
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "产品图"
    sheet.add_image(ExcelImage(image_buffer), "B1")
    workbook_buffer = io.BytesIO()
    workbook.save(workbook_buffer)
    workbook.close()

    parsed = await NativeDocumentParser().parse(
        filename="image-spec.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=workbook_buffer.getvalue(),
    )

    assert parsed.report.complete is False
    assert "xlsx_images_not_analyzed" in {
        notice.code for notice in parsed.report.notices
    }


@pytest.mark.asyncio
async def test_native_parser_rejects_excel_with_too_many_sheets() -> None:
    workbook = Workbook()
    for index in range(20):
        workbook.create_sheet(f"sheet-{index + 2}")
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    with pytest.raises(ValidationError, match="超过 20 个工作表"):
        await NativeDocumentParser().parse(
            filename="too-many-sheets.xlsx",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=buffer.getvalue(),
        )


@pytest.mark.asyncio
async def test_fake_product_import_returns_editable_governed_draft(session_factory) -> None:
    async with session_factory() as session, session.begin():
        brand = Brand(name="Import Brand", slug="import-brand")
        session.add(brand)
        await session.flush()
        brand_id = brand.id
        actor_id = brand.id

    async with session_factory() as session:
        registry = RegistryService(session)
        _item, prompt = await registry.create_item(
            brand_id=brand_id,
            registry_type=RegistryType.PROMPT_TEMPLATE,
            actor_id=actor_id,
            data=RegistryItemCreate(
                stable_key="product_registry_import",
                payload={
                    "purpose": "product_registry_import",
                    "template": "只整理输入资料并返回结构化草稿。",
                    "variables": [],
                    "output_schema_name": "ProductImportAIOutput",
                },
            ),
        )
        await registry.activate_revision(
            brand_id=brand_id,
            revision_id=prompt.id,
            actor_id=actor_id,
        )

    draft = await ProductImportService(
        session_factory,
        DevelopmentFakeLLMGateway(),
    ).create_draft(
        brand_id=brand_id,
        pasted_text="Paper Orbit NOTE01 点阵笔记本产品规格书，具体功能仍待最终确认。",
        files=[],
    )
    assert draft.payload.sku == "NOTE01"
    assert draft.stable_key == "paper-orbit-note01-ai-draft"
    assert draft.missing_fields
    assert any("不得直接" in warning for warning in draft.warnings)
    assert draft.dedup is not None
    assert draft.dedup.duplicate_status == "none"

    with pytest.raises(ConflictError, match="本次未调用 LLM"):
        await ProductImportService(
            session_factory,
            DevelopmentFakeLLMGateway(),
        ).create_draft(
            brand_id=brand_id,
            pasted_text="Paper Orbit NOTE01 点阵笔记本产品规格书，具体功能仍待最终确认。",
            files=[],
        )


@pytest.mark.asyncio
async def test_product_import_keeps_paid_draft_when_dedup_persistence_fails(
    session_factory,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async with session_factory() as session, session.begin():
        brand = Brand(name="Resilient Import Brand", slug="resilient-import-brand")
        session.add(brand)
        await session.flush()
        brand_id = brand.id
        actor_id = brand.id

    async with session_factory() as session:
        registry = RegistryService(session)
        _item, prompt = await registry.create_item(
            brand_id=brand_id,
            registry_type=RegistryType.PROMPT_TEMPLATE,
            actor_id=actor_id,
            data=RegistryItemCreate(
                stable_key="product_registry_import",
                payload={
                    "purpose": "product_registry_import",
                    "template": "只整理输入资料并返回结构化草稿。",
                    "variables": [],
                    "output_schema_name": "ProductImportAIOutput",
                },
            ),
        )
        await registry.activate_revision(
            brand_id=brand_id,
            revision_id=prompt.id,
            actor_id=actor_id,
        )

    async def fail_dedup(*_args, **_kwargs):
        raise SQLAlchemyError("database unavailable")

    monkeypatch.setattr(IngestionDedupService, "analyze_product", fail_dedup)
    draft = await ProductImportService(
        session_factory,
        DevelopmentFakeLLMGateway(),
    ).create_draft(
        brand_id=brand_id,
        pasted_text="Paper Orbit NOTE01 点阵笔记本产品规格书，具体功能仍待最终确认。",
        files=[],
    )

    assert draft.payload.sku == "NOTE01"
    assert draft.dedup is None
    assert any("重复内容检测记录暂时保存失败" in item for item in draft.warnings)
